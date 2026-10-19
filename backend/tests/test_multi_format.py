import io

from docx import Document as DocxDocument

from app.services.document_processor import extract_text


def test_extract_text_plain():
    assert extract_text("notes.txt", b"hello there") == "hello there"


def test_extract_text_csv():
    text = extract_text("data.csv", b"name,age\nAlice,30\nBob,25")
    assert "Alice, 30" in text
    assert "Bob, 25" in text


def test_extract_text_html():
    html = b"<html><body><h1>Title</h1><p>Some paragraph text.</p></body></html>"
    text = extract_text("page.html", html)
    assert "Title" in text
    assert "Some paragraph text." in text


def test_extract_text_docx():
    buf = io.BytesIO()
    doc = DocxDocument()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")
    doc.save(buf)
    text = extract_text("notes.docx", buf.getvalue())
    assert "First paragraph." in text
    assert "Second paragraph." in text
